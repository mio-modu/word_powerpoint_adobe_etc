"""IR -> DOCX 변환기.

DOCX(OOXML) 는 ECMA-376 / ISO 29500 로 공개된 표준이다. HWPX 출력이
실제 워드프로세서에서 확인되기 전까지 이 변환기가 대조군 역할을 한다.
같은 IR 에서 두 포맷이 나오므로 내용 차이가 생기면 writer 쪽 문제다.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt

from ..ir import Align, Document, Heading, PageBreak, Paragraph, Table

_ALIGN = {
    Align.LEFT: WD_ALIGN_PARAGRAPH.LEFT,
    Align.CENTER: WD_ALIGN_PARAGRAPH.CENTER,
    Align.RIGHT: WD_ALIGN_PARAGRAPH.RIGHT,
    Align.JUSTIFY: WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def write_docx(doc: Document, path: str | Path) -> Path:
    out = DocxDocument()

    style = out.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10)

    title = out.add_heading(doc.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for block in doc.blocks:
        match block:
            case Heading(text=text, level=level):
                out.add_heading(text, level=level)
            case Paragraph() as para:
                p = out.add_paragraph()
                p.alignment = _ALIGN[para.align]
                for run in para.runs:
                    r = p.add_run(run.text)
                    r.bold = run.bold
                    r.italic = run.italic
                    r.underline = run.underline
            case Table() as table:
                t = out.add_table(rows=0, cols=table.n_cols)
                t.style = "Table Grid"
                for i, row in enumerate(table.rows):
                    cells = t.add_row().cells
                    for cell, value in zip(cells, row):
                        cell.text = value
                        if table.header and i == 0:
                            for p in cell.paragraphs:
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                for r in p.runs:
                                    r.bold = True
            case PageBreak():
                out.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            case _:
                raise TypeError(f"지원하지 않는 블록입니다: {type(block).__name__}")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    out.save(str(path))
    return path
